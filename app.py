import io
import os
import hashlib
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders

import streamlit as st
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image
import numpy as np

# Page configuration
st.set_page_config(
    page_title="Generator Dokumentasi Kegiatan Siap Cetak",
    page_icon="📄",
    layout="centered"
)

st.title("📄 Generator Dokumentasi Kegiatan")
st.markdown("Aplikasi otomatis penyusun foto dokumentasi ke dalam format **Microsoft Word (.docx)** ukuran **A4 Portrait** dengan tata letak **Grid 3×3 Siap Cetak**.")

# Sidebar Configuration for Email / SMTP
with st.sidebar:
    st.header("⚙️ Pengaturan Email (Opsional)")
    st.markdown("Jika ingin hasil dokumen otomatis terkirim ke email, aktifkan opsi ini:")
    enable_email = st.checkbox("Kirim otomatis ke email", value=False)
    
    sender_email = st.text_input("Email Pengirim (Gmail):", placeholder="contoh@gmail.com")
    app_password = st.text_input("App Password Gmail:", type="password", help="Gunakan 16 karakter App Password dari akun Google Anda (bukan password biasa).")
    smtp_server = st.text_input("SMTP Server:", value="smtp.gmail.com")
    smtp_port = st.number_input("SMTP Port:", value=587)
    
    st.divider()
    st.caption("Aplikasi dikembangkan sesuai dengan aturan format cetak laporan resmi.")

# Main Form
with st.form("doc_form"):
    st.subheader("1. Informasi Kegiatan")
    col1, col2 = st.columns(2)
    with col1:
        nama_kegiatan = st.text_input("Nama Kegiatan:", placeholder="Contoh: Musyawarah Desa")
    with col2:
        tanggal_kegiatan = st.text_input("Tanggal Kegiatan:", placeholder="Contoh: Senin, 10 Oktober 2026")
        
    email_tujuan = st.text_input("Email Penerima Dokumen:", value="hitamair4@gmail.com")
    
    st.subheader("2. Unggah Foto Dokumentasi")
    uploaded_files = st.file_uploader(
        "Pilih foto-foto kegiatan (bisa pilih banyak sekaligus):",
        type=["jpg", "jpeg", "png"],
        accept_multiple_files=True
    )
    
    submitted = st.form_submit_button("🚀 Buat Dokumen Word", use_container_width=True)

def remove_black_bars(pil_img):
    """Membersihkan bilah hitam letterbox jika foto berupa screenshot/tangkapan layar"""
    arr = np.array(pil_img.convert("RGB"))
    non_black = np.where(arr.mean(axis=(1,2)) > 15)[0]
    if len(non_black) > 0 and (non_black[0] > 10 or non_black[-1] < len(arr) - 10):
        y_start, y_end = non_black[0], non_black[-1]
        return pil_img.crop((0, y_start, pil_img.width, y_end + 1))
    return pil_img

def process_photos_to_docx(nama_keg, tgl_keg, files_list):
    doc = docx.Document()
    
    # Setup halaman A4 Portrait
    for sec in doc.sections:
        sec.page_width = Inches(8.27)   # 210 mm
        sec.page_height = Inches(11.69) # 297 mm
        sec.top_margin = Inches(0.85)
        sec.bottom_margin = Inches(0.85)
        sec.left_margin = Inches(0.70)
        sec.right_margin = Inches(0.70)
        
    # Filter duplikat berdasarkan hash MD5
    unique_images = []
    seen_hashes = set()
    duplicates_count = 0
    
    for f in files_list:
        content = f.read()
        f.seek(0)
        h = hashlib.md5(content).hexdigest()
        if h in seen_hashes:
            duplicates_count += 1
            continue
        seen_hashes.add(h)
        
        im = Image.open(f).convert("RGB")
        im = remove_black_bars(im)
        unique_images.append(im)
        
    total_photos = len(unique_images)
    if total_photos == 0:
        return None, 0, 0
        
    # Pembagian per 9 foto per halaman
    photos_per_page = 9
    pages_count = (total_photos + photos_per_page - 1) // photos_per_page
    
    for page_idx in range(pages_count):
        if page_idx > 0:
            doc.add_page_break()
            
        # Header halaman
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run("DOKUMENTASI KEGIATAN")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(14)
        r_t.font.color.rgb = RGBColor(0, 0, 0)
        
        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(18)
        p_meta.paragraph_format.line_spacing = 1.25
        
        r_keg = p_meta.add_run(f"Nama Kegiatan: {nama_keg}\n")
        r_keg.font.name = "Arial"
        r_keg.font.size = Pt(11)
        r_keg.font.color.rgb = RGBColor(0, 0, 0)
        
        r_tgl = p_meta.add_run(f"Tanggal: {tgl_keg}")
        r_tgl.font.name = "Arial"
        r_tgl.font.size = Pt(11)
        r_tgl.font.color.rgb = RGBColor(0, 0, 0)
        
        # Ambil subset foto untuk halaman ini
        start_idx = page_idx * photos_per_page
        end_idx = min(start_idx + photos_per_page, total_photos)
        page_photos = unique_images[start_idx:end_idx]
        
        # Grid 3 kolom
        num_rows = (len(page_photos) + 2) // 3
        table = doc.add_table(rows=num_rows, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        col_width = Inches(2.29)
        
        for r_i in range(num_rows):
            row = table.rows[r_i]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
            
            for c_i in range(3):
                cell = row.cells[c_i]
                cell.width = col_width
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                # Padding sel konsisten
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = parse_xml(r'<w:tcMar {}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="40" w:type="dxa"/><w:right w:w="40" w:type="dxa"/></w:tcMar>'.format(nsdecls('w')))
                tcPr.append(tcMar)
                
                # Borderless
                tcBorders = parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="none"/><w:left w:val="none"/>
                        <w:bottom w:val="none"/><w:right w:val="none"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                tcPr.append(tcBorders)
                
                cell_idx = r_i * 3 + c_i
                p_c = cell.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_c.paragraph_format.space_before = Pt(2)
                p_c.paragraph_format.space_after = Pt(4)
                p_c.paragraph_format.line_spacing = 1.0
                
                if cell_idx < len(page_photos):
                    img = page_photos[cell_idx]
                    w, h = img.size
                    
                    # Jika orientasi portrait (tinggi > lebar), gunakan canvas 16:9 putih agar grid simetris seragam
                    if h > w:
                        canvas = Image.new("RGB", (1000, 563), (255, 255, 255))
                        scale = 563.0 / h
                        new_w = int(w * scale)
                        resized_p = img.resize((new_w, 563), Image.Resampling.LANCZOS)
                        canvas.paste(resized_p, ((1000 - new_w) // 2, 0))
                        target_img = canvas
                    else:
                        target_img = img
                        
                    # Simpan ke byte stream
                    img_bytes = io.BytesIO()
                    target_img.save(img_bytes, format="JPEG", quality=95)
                    img_bytes.seek(0)
                    
                    p_c.add_run().add_picture(img_bytes, width=Inches(2.20))
                    
    # Simpan hasil akhir ke memory buffer
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_buf.seek(0)
    return docx_buf, total_photos, duplicates_count

def send_docx_via_email(docx_bytes, filename, to_email, smtp_cfg, nama_keg, tgl_keg):
    msg = MIMEMultipart()
    msg['From'] = smtp_cfg['sender']
    msg['To'] = to_email
    msg['Subject'] = f"Dokumentasi Kegiatan: {nama_keg} - {tgl_keg}"
    
    body = f"""Halo,

Terlampir dokumen Word (.docx) dokumentasi kegiatan {nama_keg} yang dilaksanakan pada {tgl_keg}.
Format dokumen sudah diatur rapi ukuran A4 Portrait dengan tata letak grid 3x3 siap cetak.

Terima kasih."""
    msg.attach(MIMEText(body, 'plain'))
    
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(docx_bytes.getvalue())
    encoders.encode_base64(part)
    part.add_header('Content-Disposition', f'attachment; filename="{filename}"')
    msg.attach(part)
    
    server = smtplib.SMTP(smtp_cfg['server'], smtp_cfg['port'])
    server.starttls()
    server.login(smtp_cfg['sender'], smtp_cfg['password'])
    server.sendmail(smtp_cfg['sender'], to_email, msg.as_string())
    server.quit()

if submitted:
    if not nama_kegiatan or not tanggal_kegiatan:
        st.error("⚠️ Mohon isi Nama Kegiatan dan Tanggal Kegiatan terlebih dahulu!")
    elif not uploaded_files:
        st.error("⚠️ Mohon unggah minimal 1 foto dokumentasi!")
    else:
        with st.spinner("Sedang memproses foto dan menyusun dokumen Word..."):
            doc_stream, total_u, dups_c = process_photos_to_docx(nama_kegiatan, tanggal_kegiatan, uploaded_files)
            
            clean_name = "".join(c if c.isalnum() else "_" for c in nama_kegiatan)
            filename = f"Dokumentasi_{clean_name}.docx"
            
            st.success(f"✅ Dokumen berhasil dibuat! Total {total_u} foto dimasukkan ({dups_c} duplikat diabaikan).")
            
            # Tombol Download Langsung
            st.download_button(
                label=f"📥 Unduh {filename}",
                data=doc_stream,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
            
            # Pengiriman Email jika diaktifkan
            if enable_email:
                if not sender_email or not app_password:
                    st.warning("⚠️ Email belum terkirim: Mohon lengkapi Email Pengirim dan App Password di menu samping (Sidebar).")
                else:
                    try:
                        smtp_cfg = {
                            "sender": sender_email,
                            "password": app_password,
                            "server": smtp_server,
                            "port": smtp_port
                        }
                        send_docx_via_email(doc_stream, filename, email_tujuan, smtp_cfg, nama_kegiatan, tanggal_kegiatan)
                        st.success(f"📧 Dokumen berhasil dikirimkan sebagai lampiran ke: **{email_tujuan}**")
                    except Exception as e:
                        st.error(f"❌ Gagal mengirim email: {str(e)}")
