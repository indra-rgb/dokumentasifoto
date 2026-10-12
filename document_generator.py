import os
import io
import hashlib
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image, ImageOps
import numpy as np

def clean_and_orient_image(pil_img):
    """Koreksi orientasi EXIF dan bersihkan bilah hitam letterbox jika ada"""
    # 1. Koreksi EXIF auto-rotate
    try:
        pil_img = ImageOps.exif_transpose(pil_img)
    except Exception:
        pass
        
    pil_img = pil_img.convert("RGB")
    
    # 2. Deteksi bilah hitam letterbox (misal dari tangkapan layar smartphone)
    try:
        arr = np.array(pil_img)
        non_black = np.where(arr.mean(axis=(1,2)) > 15)[0]
        if len(non_black) > 0 and (non_black[0] > 15 or non_black[-1] < len(arr) - 15):
            y_start, y_end = non_black[0], non_black[-1]
            # Pastikan rasio crop masih masuk akal (tidak memotong lebih dari 35% tinggi)
            if (y_end - y_start) > len(arr) * 0.65:
                pil_img = pil_img.crop((0, y_start, pil_img.width, y_end + 1))
    except Exception:
        pass
        
    return pil_img

def calculate_grid_config(col_count):
    """Menghitung konfigurasi kolom, lebar sel, dan kapasitas per halaman A4"""
    # Lebar cetak efektif A4 (8.27 in - 1.4 in margin) = 6.87 in (17.45 cm)
    printable_width = 6.87
    
    if col_count == 2:
        # 2 Kolom x 2 Baris = 4 foto/halaman
        per_page = 4
        col_width = Inches(printable_width / 2.0)
        img_width = Inches(3.25)
    elif col_count == 4:
        # 4 Kolom x 3 Baris = 12 foto/halaman
        per_page = 12
        col_width = Inches(printable_width / 4.0)
        img_width = Inches(1.62)
    elif col_count == 6:
        # 6 Kolom x 3 Baris = 18 foto/halaman
        per_page = 18
        col_width = Inches(printable_width / 6.0)
        img_width = Inches(1.08)
    else:
        # Default: 3 Kolom x 3 Baris = 9 foto/halaman
        col_count = 3
        per_page = 9
        col_width = Inches(printable_width / 3.0)
        img_width = Inches(2.20)
        
    return {
        "col_count": col_count,
        "per_page": per_page,
        "col_width": col_width,
        "img_width": img_width
    }

def format_photo_frame(img, mode="proporsional"):
    """
    Menyesuaikan foto agar seragam pada grid dokumen Word tanpa distorsi.
    - 'proporsional': Menjaga rasio asli 100%. Jika portrait (tinggi > lebar),
      foto diletakkan di kanvas 16:9 putih sehingga tinggi baris tetap seragam.
    - 'isi_bingkai': Center-crop proporsional ke 16:9 standar landscape.
    """
    w, h = img.size
    target_aspect = 16.0 / 9.0  # 1.777
    current_aspect = w / float(h)
    
    if mode == "isi_bingkai":
        # Crop bagian tengah ke rasio 16:9
        if current_aspect > target_aspect:
            # Terlalu lebar, crop sisi kiri dan kanan
            new_w = int(h * target_aspect)
            left = (w - new_w) // 2
            return img.crop((left, 0, left + new_w, h))
        else:
            # Terlalu tinggi, crop sisi atas dan bawah
            new_h = int(w / target_aspect)
            top = (h - new_h) // 2
            return img.crop((0, top, w, top + new_h))
    else:
        # Mode Proporsional (Default)
        if h > w:
            # Foto portrait: letakkan di kanvas 16:9 berlatar putih bersih
            canvas_w = int(h * target_aspect)
            canvas = Image.new("RGB", (canvas_w, h), (255, 255, 255))
            left = (canvas_w - w) // 2
            canvas.paste(img, (left, 0))
            return canvas
        return img

def generate_docx_document(nama_kegiatan, tanggal_kegiatan, lokasi, images_list, col_count=3, orientation_mode="proporsional", show_captions=False):
    """Menghasilkan berkas Word .docx siap cetak A4 Portrait"""
    doc = docx.Document()
    
    # 1. Pengaturan Halaman A4 Portrait
    for sec in doc.sections:
        sec.page_width = Inches(8.27)   # 210 mm
        sec.page_height = Inches(11.69) # 297 mm
        sec.top_margin = Inches(0.85)
        sec.bottom_margin = Inches(0.85)
        sec.left_margin = Inches(0.70)
        sec.right_margin = Inches(0.70)
        
    cfg = calculate_grid_config(col_count)
    cols = cfg["col_count"]
    per_page = cfg["per_page"]
    col_w = cfg["col_width"]
    img_w = cfg["img_width"]
    
    total_photos = len(images_list)
    pages_count = max(1, (total_photos + per_page - 1) // per_page)
    
    for page_idx in range(pages_count):
        if page_idx > 0:
            doc.add_page_break()
            
        # Header Halaman
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run("DOKUMENTASI KEGIATAN")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(14)
        r_t.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        
        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(16)
        p_meta.paragraph_format.line_spacing = 1.25
        
        r_k = p_meta.add_run(f"Nama Kegiatan: {nama_kegiatan}\n")
        r_k.font.name = "Arial"
        r_k.font.size = Pt(11)
        r_k.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        
        tgl_text = f"Tanggal: {tanggal_kegiatan}"
        if lokasi:
            tgl_text += f"\nLokasi: {lokasi}"
        r_d = p_meta.add_run(tgl_text)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        
        # Ambil subset foto halaman ini
        start_idx = page_idx * per_page
        end_idx = min(start_idx + per_page, total_photos)
        cur_photos = images_list[start_idx:end_idx]
        
        num_rows = max(1, (len(cur_photos) + cols - 1) // cols)
        table = doc.add_table(rows=num_rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_i in range(num_rows):
            row = table.rows[r_i]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
            
            for c_i in range(cols):
                cell = row.cells[c_i]
                cell.width = col_w
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                # Padding sel
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = parse_xml(r'<w:tcMar {}><w:top w:w="60" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="30" w:type="dxa"/><w:right w:w="30" w:type="dxa"/></w:tcMar>'.format(nsdecls('w')))
                tcPr.append(tcMar)
                
                # Borderless
                tcBorders = parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="none"/><w:left w:val="none"/>
                        <w:bottom w:val="none"/><w:right w:val="none"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                tcPr.append(tcBorders)
                
                p_c = cell.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_c.paragraph_format.space_before = Pt(2)
                p_c.paragraph_format.space_after = Pt(4)
                p_c.paragraph_format.line_spacing = 1.0
                
                cell_photo_idx = r_i * cols + c_i
                if cell_photo_idx < len(cur_photos):
                    item = cur_photos[cell_photo_idx]
                    raw_img = item["image"]
                    caption_text = item.get("caption", f"Foto {start_idx + cell_photo_idx + 1}")
                    
                    processed_img = format_photo_frame(raw_img, mode=orientation_mode)
                    
                    stream = io.BytesIO()
                    processed_img.save(stream, format="JPEG", quality=95)
                    stream.seek(0)
                    
                    p_c.add_run().add_picture(stream, width=img_w)
                    
                    if show_captions:
                        p_cap = cell.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(1)
                        p_cap.paragraph_format.space_after = Pt(2)
                        r_cap = p_cap.add_run(caption_text)
                        r_cap.font.name = "Arial"
                        r_cap.font.size = Pt(8.5)
                        r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                        
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_io.seek(0)
    return docx_io

def convert_docx_to_pdf(docx_path, output_dir):
    """Mengonversi dokumen Word ke PDF menggunakan LibreOffice"""
    try:
        cmd = [
            "libreoffice",
            "--headless",
            "--nologo",
            "--norestore",
            "--convert-to", "pdf",
            docx_path,
            "--outdir", output_dir
        ]
        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
        base = os.path.splitext(os.path.basename(docx_path))[0]
        pdf_path = os.path.join(output_dir, f"{base}.pdf")
        if os.path.exists(pdf_path):
            return pdf_path
    except Exception as e:
        print(f"Konversi PDF gagal: {e}")
    return None
